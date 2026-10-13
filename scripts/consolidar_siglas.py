"""Consolida fontes locais sem confundir abreviaturas com identidades confirmadas."""

import argparse
import csv
import hashlib
import io
import json
import re
import unicodedata
from collections import defaultdict
from pathlib import Path


def key(text):
    text = unicodedata.normalize("NFKD", str(text or ""))
    return "".join(c.lower() for c in text if c.isalnum() and not unicodedata.combining(c))


def token(text):
    return re.sub(r"[^A-Za-z0-9]", "", unicodedata.normalize("NFKD", str(text or "")))[:20]


def full_name(name, acronym):
    """Remove apenas uma sigla concordante isolada nos extremos do nome."""
    parts = re.split(r"\s[-–—]\s|[()]", name)
    while parts and not parts[-1].strip():
        parts.pop()
    if len(parts) > 1 and key(parts[0]) == key(acronym):
        parts.pop(0)
    elif len(parts) > 1 and key(parts[-1]) == key(acronym):
        parts.pop()
    return key(" ".join(parts))


def consolidate(docs, complement, output, *, decisions=None):
    docs, complement, output = Path(docs), Path(complement), Path(output)
    if output.exists():
        raise FileExistsError("A saída já existe; escolha um ficheiro de proposta novo para conservar decisões anteriores.")
    rows = {}
    evidence = []

    def add(name, acronym, code, side, source, state, origin):
        name = re.sub(r"\s+", " ", str(name or "")).strip()
        if not name or name == "0":
            return
        identity = (key(name), code, side, token(acronym))
        if identity in rows:
            row = rows[identity]
            row["fonte"] = " | ".join(sorted(set(row["fonte"].split(" | ") + [source])))
            return
        rows[identity] = {"nome": name, "sigla": token(acronym), "codigo_dgert": code,
                          "lado": side, "estado": state, "origem_sigla": origin, "fonte": source,
                          "responsavel": "", "data": ""}

    def source(path):
        evidence.append({"ficheiro": path.name,
                         "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})

    source(complement)
    with complement.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle, delimiter=";"):
            confirmed = row["origem_sigla"].split("+")[0] in {"registo", "derivada"}
            add(row["denominacao"], row["sigla"], row["codigo_dgert"], row["lado"],
                "AppCCT/" + complement.name, "documentada" if confirmed else "por_confirmar",
                row["origem_sigla"])
    by_code = {}
    for row in rows.values():
        if row["codigo_dgert"]:
            by_code.setdefault((row["codigo_dgert"], row["lado"]), []).append(row)
    import openpyxl
    for path in sorted(docs.glob("*.xlsx")):
        source(path)
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for title, side in (("Organizações sindicais", "trabalhadores"),
                            ("Organizações de empregadores", "empregadores")):
            if title not in workbook.sheetnames:
                continue
            values = workbook[title].iter_rows(values_only=True)
            header = next(values, ())
            for number, values_row in enumerate(values, 2):
                item = dict(zip(header, values_row))
                code = str(item.get("Código Identificador da Organização") or "")
                candidates = by_code.get((code, side), [])
                known = candidates[0] if len({row["sigla"] for row in candidates}) == 1 else None
                acronym = known["sigla"] if known else item.get("Acrónimo") or ""
                state = known["estado"] if known else "documentada" if acronym else "por_confirmar"
                add(item.get("Denominação da Organização"), acronym, code, side,
                    f"{path.name}/{title}/linha {number}", state,
                    known["origem_sigla"] if known else "registo")
        # Os outorgantes dos índices ajudam a ligar variantes de denominação.
        for sheet in workbook:
            values = sheet.iter_rows(values_only=True)
            header = next(values, ())
            if "Outorgantes" not in header:
                continue
            index = list(header).index("Outorgantes")
            for number, values_row in enumerate(values, 2):
                for name in str(values_row[index] or "").split(";"):
                    if name.strip():
                        side = "trabalhadores" if re.search(r"sindic|federa.*trabalh", name, re.I) else "empregadores"
                        add(name, "", "", side, f"{path.name}/{sheet.title}/linha {number}",
                            "por_confirmar", "indice")
        workbook.close()
    # Esta lista usa caracteres literais '\\r\\n' em algumas linhas.
    for path in sorted(docs.glob("*.csv")):
        source(path)
        text = path.read_text(encoding="utf-8-sig").replace("\\r\\n", "\n")
        for number, row in enumerate(csv.DictReader(io.StringIO(text), delimiter=";"), 2):
            for column, side in (("Empregadores", "empregadores"), ("Sindicatos", "trabalhadores")):
                add(row.get(column), "", "", side, f"{path.name}/linha {number}/{column}",
                    "por_confirmar", "lista_equipa")
    aliases = defaultdict(dict)
    full_names = defaultdict(dict)
    for row in rows.values():
        if row["estado"] == "documentada" and row["sigla"]:
            identity = ".".join(row["codigo_dgert"].split(".")[:2]) or key(row["nome"])
            for alias in (key(row["nome"]), key(row["sigla"])):
                aliases[(alias, row["lado"])][identity] = row
            full_names[(full_name(row["nome"], row["sigla"]), row["lado"],
                        key(row["sigla"]))][identity] = row
    for row in rows.values():
        if row["sigla"]:
            continue
        candidates = aliases.get((key(row["nome"]), row["lado"]), {})
        if len(candidates) == 1:
            candidate = next(iter(candidates.values()))
            row.update(sigla=candidate["sigla"], codigo_dgert=candidate["codigo_dgert"],
                       estado="documentada", origem_sigla="correspondencia_exata")
        else:
            # Acrónimo literal no índice: evidência do próprio nome, conservada.
            pieces = re.split(r"\s[-–—]\s|[()]", row["nome"])
            literal = [piece.strip() for piece in pieces
                       if re.fullmatch(r"[A-Z][A-Z0-9.]{1,19}", piece.strip())]
            if len(literal) == 1 and row["origem_sigla"] == "indice":
                candidates = full_names.get((full_name(row["nome"], literal[0]),
                                             row["lado"], key(literal[0])), {})
                if len(candidates) == 1:
                    candidate = next(iter(candidates.values()))
                    row.update(sigla=candidate["sigla"], codigo_dgert=candidate["codigo_dgert"],
                               estado="documentada", origem_sigla="nome_completo_exato")
                else:
                    row.update(sigla=token(literal[0]), estado="documentada", origem_sigla="indice_literal")
                continue
            # Proposta legível, nunca usada automaticamente na nomeação.
            row["sigla"] = token("".join(word.capitalize() for word in re.findall(r"\w+", row["nome"])
                                        if key(word) not in {"de", "da", "do", "e", "sa", "lda"}))
    if decisions is not None:
        decisions = Path(decisions)
        source(decisions)
        confirmed = {}
        confirmed_names = {}
        with decisions.open(encoding="utf-8-sig", newline="") as handle:
            for decision in csv.DictReader(handle, delimiter=";"):
                if decision.get("estado") != "confirmada":
                    continue
                if not all(decision.get(field, "").strip() for field in ("nome", "sigla", "fonte", "responsavel", "data")):
                    raise ValueError("Decisão confirmada sem nome, sigla, fonte, responsável ou data; complete o CSV antes de consolidar.")
                if token(decision["sigla"]) != decision["sigla"]:
                    raise ValueError("Sigla confirmada inválida; conserve e reveja a decisão.")
                name = key(decision["nome"])
                acronym = decision["sigla"]
                previous_acronym = confirmed_names.setdefault(name, acronym)
                if previous_acronym != acronym:
                    raise ValueError(
                        "Decisões confirmadas contraditórias para a mesma denominação; "
                        "conserve o CSV e resolva o conflito antes de consolidar."
                    )
                identity = (name, decision.get("codigo_dgert", ""), decision.get("lado", ""), acronym)
                current = {field: decision.get(field, "") for field in (
                    "nome", "sigla", "codigo_dgert", "lado", "estado", "origem_sigla", "fonte", "responsavel", "data")}
                previous = confirmed.get(identity)
                if previous is not None:
                    # Duplicados concordantes conservam toda a evidência humana.
                    for field in ("origem_sigla", "fonte", "responsavel", "data"):
                        previous[field] = " | ".join(sorted(set(
                            value for value in (*previous[field].split(" | "), *current[field].split(" | ")) if value
                        )))
                else:
                    confirmed[identity] = current
        # Só depois de validar todas as decisões é seguro substituir propostas.
        for identity, decision in confirmed.items():
            for old in list(rows):
                if old[:3] == identity[:3]:
                    del rows[old]
            rows[identity] = decision
    claims = defaultdict(set)
    for row in rows.values():
        identity = ".".join(row["codigo_dgert"].split(".")[:2]) or key(row["nome"])
        if row["estado"] in {"documentada", "confirmada"}:
            claims[row["sigla"].casefold()].add(identity)
    for row in rows.values():
        if row["estado"] in {"documentada", "confirmada"} and len(claims[row["sigla"].casefold()]) > 1:
            if row["estado"] == "confirmada":
                raise ValueError("Colisão com uma sigla confirmada; conserve a decisão e reveja as fontes antes de gerar a proposta.")
            row["estado"] = "ambigua"
    name_claims = defaultdict(set)
    for row in rows.values():
        if row["estado"] in {"documentada", "confirmada"}:
            name_claims[key(row["nome"])].add(row["sigla"].casefold())
    for row in rows.values():
        if row["estado"] == "documentada" and len(name_claims[key(row["nome"])]) > 1:
            row["estado"] = "ambigua"
    ordered = sorted(rows.values(), key=lambda row: (row["lado"], key(row["nome"]), row["codigo_dgert"]))
    if not rows:
        raise ValueError("Nenhuma denominação válida nas fontes; não gerar um vocabulário vazio.")
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("x", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(ordered[0]), delimiter=";", lineterminator="\n")
        writer.writeheader()
        writer.writerows(ordered)
    manifest = {"fontes": evidence, "linhas": len(ordered), "estados": {
        state: sum(row["estado"] == state for row in ordered)
        for state in sorted({row["estado"] for row in ordered})},
        "sha256_saida": hashlib.sha256(output.read_bytes()).hexdigest()}
    output.with_suffix(".manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--docs", type=Path, default=Path("docs"))
    parser.add_argument("--complemento", type=Path, default=Path("vocabularios/siglas_organizacoes.csv"))
    parser.add_argument("--saida", type=Path, default=Path("siglas.proposta.csv"))
    parser.add_argument("--decisoes", type=Path, help="CSV ativo cujas decisões confirmadas serão preservadas")
    args = parser.parse_args()
    print(json.dumps(consolidate(args.docs, args.complemento, args.saida, decisions=args.decisoes), ensure_ascii=False, indent=2))
